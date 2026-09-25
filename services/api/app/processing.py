from __future__ import annotations

import asyncio
import hashlib
import io
from datetime import datetime, timezone
from typing import Any

from fastapi import HTTPException

from .config import get_settings
from .extraction import TextDocument, classify_document, extract_fields
from .schemas import ProcessResult
from .supabase import SupabaseGateway


def _extract_docx(content: bytes) -> tuple[str, str, float]:
    from docx import Document

    document = Document(io.BytesIO(content))
    parts = [paragraph.text for paragraph in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(value for value in parts if value.strip()), "docx_text", .98


def _ocr_image(content: bytes) -> tuple[str, str, float]:
    import cv2
    import numpy as np
    import pytesseract

    settings = get_settings()
    if settings.tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = settings.tesseract_cmd
    image = cv2.imdecode(np.frombuffer(content, dtype=np.uint8), cv2.IMREAD_COLOR)
    if image is None:
        raise ValueError("The image could not be decoded.")
    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    gray = cv2.fastNlMeansDenoising(gray, None, 10, 7, 21)
    binary = cv2.adaptiveThreshold(gray, 255, cv2.ADAPTIVE_THRESH_GAUSSIAN_C, cv2.THRESH_BINARY, 31, 15)
    data = pytesseract.image_to_data(binary, config="--oem 3 --psm 6", output_type=pytesseract.Output.DICT)
    words: list[str] = []
    confidence_values: list[float] = []
    current_line: tuple[int, int, int] | None = None
    for index, value in enumerate(data["text"]):
        word = value.strip()
        if not word:
            continue
        line = (data["block_num"][index], data["par_num"][index], data["line_num"][index])
        if current_line is not None and line != current_line:
            words.append("\n")
        words.append(word)
        current_line = line
        try:
            confidence = float(data["conf"][index])
            if confidence >= 0:
                confidence_values.append(confidence / 100)
        except (TypeError, ValueError):
            pass
    text = " ".join(words).replace(" \n ", "\n")
    confidence = sum(confidence_values) / len(confidence_values) if confidence_values else .55
    return text, "ocr", min(.95, max(.40, confidence))


def _extract_pdf(content: bytes) -> tuple[str, str, float]:
    import pymupdf

    with pymupdf.open(stream=content, filetype="pdf") as document:
        text = "\n".join(page.get_text("text") for page in document).strip()
        if len(text) >= 40:
            return text, "pdf_text", .99
        pages: list[str] = []
        confidences: list[float] = []
        for page in document:
            pixmap = page.get_pixmap(matrix=pymupdf.Matrix(2, 2), alpha=False)
            page_text, _, confidence = _ocr_image(pixmap.tobytes("png"))
            pages.append(page_text)
            confidences.append(confidence)
        return "\n".join(pages), "ocr", (sum(confidences) / len(confidences) if confidences else .55)


def extract_document_text(content: bytes, mime_type: str) -> tuple[str, str, float]:
    if mime_type == "application/pdf":
        return _extract_pdf(content)
    if mime_type == "application/vnd.openxmlformats-officedocument.wordprocessingml.document":
        return _extract_docx(content)
    if mime_type in {"image/jpeg", "image/png"}:
        return _ocr_image(content)
    raise ValueError(f"Unsupported document type: {mime_type}")


async def process_claim(
    gateway: SupabaseGateway,
    claim: dict[str, Any],
    actor_id: str,
) -> ProcessResult:
    claim_id = claim["id"]
    reference = claim["reference"]
    now = datetime.now(timezone.utc).isoformat()
    jobs = await gateway.rest(
        "POST",
        "claim_processing_jobs",
        service=True,
        json={"claim_id": claim_id, "created_by": actor_id, "status": "RUNNING", "attempt_count": 1, "started_at": now},
        prefer="return=representation",
    )
    job_id = jobs[0]["id"]
    try:
        await gateway.rest("PATCH", "claims", service=True, params={"id": f"eq.{claim_id}"}, json={"status": "PROCESSING"})
        await gateway.rest(
            "POST",
            "audit_events",
            service=True,
            json={"claim_id": claim_id, "actor_id": actor_id, "event_type": "PROCESSING_STARTED"},
        )
    except Exception as exc:
        error = str(exc)[:1000]
        await gateway.rest("PATCH", "claims", service=True, params={"id": f"eq.{claim_id}"}, json={"status": "PROCESSING_FAILED"})
        await gateway.rest(
            "PATCH",
            "claim_processing_jobs",
            service=True,
            params={"id": f"eq.{job_id}"},
            json={"status": "FAILED", "last_error": error, "finished_at": datetime.now(timezone.utc).isoformat()},
        )
        raise HTTPException(status_code=422, detail=error) from exc
    try:
        documents = await gateway.rest(
            "GET",
            "claim_documents",
            service=True,
            params={"claim_id": f"eq.{claim_id}", "select": "*", "order": "created_at.asc"},
        )
        await gateway.rest("DELETE", "claim_extracted_fields", service=True, params={"claim_id": f"eq.{claim_id}"})
        hashes: dict[str, str] = {}
        field_count = 0
        warning_count = 0
        for document in documents:
            content = await gateway.download(document["storage_path"])
            digest = hashlib.sha256(content).hexdigest()
            duplicate_of = hashes.get(digest)
            hashes.setdefault(digest, document["id"])
            try:
                text, method, confidence = await asyncio.to_thread(extract_document_text, content, document["mime_type"])
                if not text.strip():
                    raise ValueError("No readable text was found. Upload a clearer image or enter the values manually.")
                document_type, _, payable = classify_document(text)
                fields, amount, currency, warnings = extract_fields(TextDocument(document["id"], text, method, confidence))
                warning_count += warnings + (1 if duplicate_of else 0)
                field_rows = [
                    {
                        "claim_id": claim_id,
                        "document_id": field.document_id,
                        "field_name": field.field_name,
                        "raw_value": field.raw_value,
                        "normalized_value": field.normalized_value,
                        "confidence": field.confidence,
                        "extraction_method": field.extraction_method,
                        "page_number": field.page_number,
                    }
                    for field in fields
                ]
                if field_rows:
                    await gateway.rest("POST", "claim_extracted_fields", service=True, json=field_rows)
                field_count += len(field_rows)
                amount_field = next((field for field in fields if field.field_name == "claimed_amount"), None)
                status = "COMPLETED" if amount and currency else "NEEDS_CONFIRMATION"
                notes = None if amount else "No reliable payable total was found; please enter the document amount."
                await gateway.rest(
                    "PATCH",
                    "claim_documents",
                    service=True,
                    params={"id": f"eq.{document['id']}"},
                    json={
                        "document_type": document_type,
                        "extracted_amount": amount,
                        "extracted_currency": currency,
                        "amount_confidence": amount_field.confidence if amount_field else None,
                        "include_in_total": bool(payable and amount and not duplicate_of),
                        "duplicate_of": duplicate_of,
                        "extraction_status": status,
                        "extraction_notes": notes,
                        "extracted_at": now,
                    },
                )
            except Exception as exc:
                warning_count += 1
                await gateway.rest(
                    "PATCH",
                    "claim_documents",
                    service=True,
                    params={"id": f"eq.{document['id']}"},
                    json={"extraction_status": "FAILED", "extraction_notes": str(exc)[:500], "extracted_at": now},
                )

        await gateway.rest(
            "PATCH",
            "claims",
            service=True,
            params={"id": f"eq.{claim_id}"},
            json={"status": "UPLOADED", "warning_count": warning_count},
        )
        await gateway.rest(
            "PATCH",
            "claim_processing_jobs",
            service=True,
            params={"id": f"eq.{job_id}"},
            json={"status": "COMPLETED", "finished_at": datetime.now(timezone.utc).isoformat()},
        )
        await gateway.rest(
            "POST",
            "audit_events",
            service=True,
            json={"claim_id": claim_id, "actor_id": actor_id, "event_type": "PROCESSING_COMPLETED", "metadata": {"documents": len(documents), "fields": field_count}},
        )
        return ProcessResult(claim_reference=reference, status="UPLOADED", document_count=len(documents), extracted_field_count=field_count)
    except Exception as exc:
        error = str(exc)[:1000]
        await gateway.rest("PATCH", "claims", service=True, params={"id": f"eq.{claim_id}"}, json={"status": "PROCESSING_FAILED"})
        await gateway.rest(
            "PATCH",
            "claim_processing_jobs",
            service=True,
            params={"id": f"eq.{job_id}"},
            json={"status": "FAILED", "last_error": error, "finished_at": datetime.now(timezone.utc).isoformat()},
        )
        await gateway.rest(
            "POST",
            "audit_events",
            service=True,
            json={"claim_id": claim_id, "actor_id": actor_id, "event_type": "PROCESSING_FAILED", "metadata": {"error": error}},
        )
        raise HTTPException(status_code=422, detail=error) from exc
